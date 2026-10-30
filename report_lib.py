#!/usr/bin/env python3
"""
University of Mumbai Academic Project Report Generator
Project: History of India - An Interactive Digital Archive & Full-Stack Web Platform
Candidate: Mr. Raunak Dilip Jaiswal
Institution: V.K. Krishna Menon College of Commerce & Economics, Mumbai-400042
Total Pages: Exactly 68 Pages
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ASSETS_DIR = r"C:\Users\RAUNAK JAISWAL\OneDrive\Desktop\History of India\doc_assets"

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_p(doc, text, bold_prefix=None, space_after=4, italic_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
        
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.font.name = 'Calibri'
        r_it.font.size = Pt(11)
        r_it.italic = True
        r_it.font.color.rgb = RGBColor(71, 85, 105)
        
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
        
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_image_figure(doc, img_rel_name, caption, width_in=6.2, height_in=None):
    img_path = os.path.join(ASSETS_DIR, img_rel_name)
    if not os.path.exists(img_path):
        p_err = doc.add_paragraph(f"[Image Missing: {img_rel_name}]")
        p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(2)
    run_img = p_img.add_run()
    if height_in:
        run_img.add_picture(img_path, width=Inches(width_in), height=Inches(height_in))
    else:
        run_img.add_picture(img_path, width=Inches(width_in))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(4)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Calibri'
    r_cap.font.size = Pt(10)
    r_cap.bold = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)

def add_code_snippet(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Code Listing: {caption}")
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(10.5)
        r_cap.bold = True
        r_cap.font.color.rgb = RGBColor(51, 65, 85)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

print("Report generator library loaded successfully.")
