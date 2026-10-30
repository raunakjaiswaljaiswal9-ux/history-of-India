import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_lib import (
    add_heading_1, add_heading_2, add_heading_3, add_p, add_bullet, 
    add_image_figure, set_cell_background, set_cell_margins
)

def build_p56(doc):
    # Page 56: Chapter 12: Database Design (Part 1)
    add_heading_1(doc, "12. DATABASE DESIGN & RELATIONAL SCHEMA")
    add_p(doc, 
          "The data architecture of the 'History of India' digital archive is engineered to support structured, high-speed, "
          "relational queries across civilizational eras, historical dynasties, curated museum artifacts, educational quizzes, "
          "and user bookmarks. The persistence layer operates across dual modalities: localized JSON data repositories for "
          "instant client-side loading, and a normalized SQLite3 relational schema for persistent server-side archiving.", space_after=6)

    add_heading_2(doc, "12.1 Relational Architecture & Normalization")
    add_p(doc, 
          "The physical schema is normalized to Third Normal Form (3NF) to eliminate data redundancy and preserve referential "
          "integrity through strictly enforced foreign keys across five primary entities:", space_after=6)
    add_bullet(doc, "Master catalog of chronological civilizational periods, date boundaries, summaries, and thematic accent hex colors.", bold_prefix="1. Table `eras`: ")
    add_bullet(doc, "Imperial ruling houses and polities linked via foreign key (`era_id`) to their respective epoch.", bold_prefix="2. Table `dynasties`: ")
    add_bullet(doc, "Curated museum relics, sculptures, and manuscripts linked via (`era_id`), storing physical media, provenance, and geographic coordinates.", bold_prefix="3. Table `artifacts`: ")
    add_bullet(doc, "Pedagogical assessment item bank containing multiple-choice questions, options, and explanations linked via (`era_id`).", bold_prefix="4. Table `quiz_questions`: ")
    add_bullet(doc, "User-selected favorite artifacts linked via (`artifact_id`) with persistent timestamps.", bold_prefix="5. Table `bookmarks`: ")

def build_p57(doc):
    # Page 57: Chapter 12: Database Design (Part 2 - Data Dictionary)
    add_heading_2(doc, "12.2 Data Dictionary & Field Specifications")
    tbl_dict = doc.add_table(rows=1, cols=4)
    tbl_dict.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl_dict.rows[0].cells
    hdr[0].width = Inches(1.5)
    hdr[1].width = Inches(1.4)
    hdr[2].width = Inches(1.1)
    hdr[3].width = Inches(2.6)
    for i, h in enumerate(["Field Name", "Data Type", "Constraint", "Description & Operational Function"]):
        set_cell_background(hdr[i], "1E293B")
        set_cell_margins(hdr[i], top=45, bottom=45, left=80, right=80)
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    dict_rows = [
        ("era_id", "VARCHAR(32)", "PRIMARY KEY", "Unique slug identifier for the civilizational era (e.g. 'ancient')"),
        ("name", "VARCHAR(128)", "NOT NULL", "Official historical designation of the era or imperial dynasty"),
        ("date_range", "VARCHAR(64)", "NOT NULL", "Chronological epoch span (e.g. 'c. 2500 BCE – 322 BCE')"),
        ("artifact_id", "VARCHAR(64)", "PRIMARY KEY", "Unique identifier for curated museum relic (e.g. 'dancing-girl')"),
        ("medium", "VARCHAR(128)", "NOT NULL", "Archaeological material (e.g. 'Bronze lost-wax casting')"),
        ("origin", "VARCHAR(128)", "NOT NULL", "Excavation site or geographical place of origin"),
        ("image_url", "VARCHAR(256)", "NOT NULL", "Local relative file path to high-resolution asset"),
        ("latitude", "DECIMAL(8,4)", "NOT NULL", "Calibrated latitude coordinate for cartographic mapping"),
        ("longitude", "DECIMAL(8,4)", "NOT NULL", "Calibrated longitude coordinate for cartographic mapping"),
        ("question_id", "INTEGER", "PRIMARY KEY", "Auto-incrementing unique identifier for quiz question"),
        ("correct_idx", "TINYINT", "NOT NULL", "Zero-based index of the mathematically correct option (0-3)")
    ]
    for fn, dt, cn, ds in dict_rows:
        row = tbl_dict.add_row().cells
        row[0].width = Inches(1.5)
        row[1].width = Inches(1.4)
        row[2].width = Inches(1.1)
        row[3].width = Inches(2.6)
        bg = "F8FAFC" if (len(tbl_dict.rows) % 2 == 0) else "FFFFFF"
        for i, val in enumerate([fn, dt, cn, ds]):
            set_cell_background(row[i], bg)
            set_cell_margins(row[i], top=25, bottom=25, left=60, right=60)
            p = row[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)

def build_p58(doc):
    # Page 58: Chapter 12: Database Design (Part 3 - Relational Schema Diagram)
    add_heading_2(doc, "Figure 12.1: Relational Database Schema Architecture")
    add_image_figure(doc, "database_schema_diagram.png", 
                     "Figure 12.1: Physical Relational Database Schema & Foreign Key Constraints", width_in=6.2)
    add_p(doc, 
          "Figure 12.1 illustrates the physical entity structures, column data types, primary key definitions, and foreign "
          "key constraints linking `eras`, `dynasties`, `artifacts`, `quiz_questions`, and `bookmarks` within the relational store.", space_after=4)

def build_p59(doc):
    # Page 59: Chapter 13: Screenshots - Intro & Screenshot 1
    add_heading_1(doc, "13. SCREENSHOTS OF THE APPLICATION")
    add_p(doc, 
          "The 'History of India' digital archive provides an immersive, responsive user experience adhering to the "
          "Sony Darkroom design philosophy. The following high-resolution production captures illustrate the operational "
          "capabilities, visual elegance, and interactive modules of the live web application:", space_after=4)
    add_heading_2(doc, "1. Hero Portal & Interactive Navigation Canvas")
    add_image_figure(doc, "ss1_hero_portal.png", 
                     "Figure 13.1: Hero Portal displaying Sony Darkroom Theme, Emblem & Navigation Suite", width_in=5.8, height_in=3.6)
    add_p(doc, 
          "The Hero Portal introduces the user to 5,000 years of Indian civilizational history, featuring a high-contrast dark "
          "aesthetic, real-time search input, Tanpura synthesizer toggle, and view mode switcher.", space_after=4)

def build_p60(doc):
    # Page 60: Chapter 13: Screenshots 2 & 3
    add_heading_2(doc, "2. Dual-Mode Horizontal Timeline Carousel")
    add_image_figure(doc, "ss2_horizontal_timeline.png", 
                     "Figure 13.2: Horizontal Chronological Timeline Carousel with Wheel Scrolling", width_in=5.6, height_in=2.7)
    add_heading_2(doc, "3. Vertical Chronological Feed")
    add_image_figure(doc, "ss3_vertical_feed.png", 
                     "Figure 13.3: Vertical Chronological Feed with Illuminated Milestone Cards", width_in=5.6, height_in=2.7)

def build_p61(doc):
    # Page 61: Chapter 13: Screenshots 4 & 5
    add_heading_2(doc, "4. Curated Museum Artifact Gallery")
    add_image_figure(doc, "ss4_artifact_gallery.png", 
                     "Figure 13.4: Museum Artifact Exhibition Grid with Category Filter Chips", width_in=5.6, height_in=2.7)
    add_heading_2(doc, "5. High-Resolution Lightbox Modal")
    add_image_figure(doc, "ss5_lightbox_modal.png", 
                     "Figure 13.5: Lightbox Modal showing High-Res Sculpture Inspection & Metadata", width_in=5.6, height_in=2.7)

def build_p62(doc):
    # Page 62: Chapter 13: Screenshot 6
    add_heading_2(doc, "6. Imperial Interactive Atlas & Territory Geographic Map")
    add_image_figure(doc, "ss6_imperial_atlas.png", 
                     "Figure 13.6: Imperial Interactive Atlas displaying Subcontinent Vector Map & Plotted Dynasty Pins", width_in=5.8, height_in=3.8)
    add_p(doc, 
          "The Imperial Atlas renders an interactive vector map of historical India. Users can click on archaeological sites "
          "and dynastic capitals (such as Mohenjo-daro, Pataliputra, and Thanjavur) to view comprehensive historical dossiers.", space_after=4)

def build_p63(doc):
    # Page 63: Chapter 13: Screenshots 7 & 8
    add_heading_2(doc, "7. Historical Knowledge Quiz Engine")
    add_image_figure(doc, "ss7_knowledge_quiz.png", 
                     "Figure 13.7: Interactive Self-Assessment Quiz Suite with Real-Time Feedback", width_in=5.6, height_in=2.7)
    add_heading_2(doc, "8. Real-Time Multi-Index Search Terminal")
    add_image_figure(doc, "ss8_instant_search.png", 
                     "Figure 13.8: Instant Substring Search Terminal Filtering Relics & Dynasties", width_in=5.6, height_in=2.7)

def build_p64(doc):
    # Page 64: Chapter 13: Screenshots 9 & 10
    add_heading_2(doc, "9. Ambient Tanpura Web Audio Synthesizer Console")
    add_image_figure(doc, "ss9_audio_atmosphere.png", 
                     "Figure 13.9: Ambient Tanpura Synthesizer Controls (Oscillator Gain & Frequency Nodes)", width_in=5.6, height_in=2.7)
    add_heading_2(doc, "10. Python HTTP Server & Localhost REST Service Console")
    add_image_figure(doc, "ss10_python_server.png", 
                     "Figure 13.10: Python 3.12 HTTP Server Dispatching REST Endpoints & Serving Assets", width_in=5.6, height_in=2.7)

def build_p65(doc):
    # Page 65: Chapter 14: Future Scope
    add_heading_1(doc, "14. FUTURE SCOPE")
    add_p(doc, 
          "The 'History of India' digital archive establishes a robust foundation for interactive cultural heritage informatics. "
          "Future development will focus on scaling the system from a standalone web terminal into an institutional, "
          "decentralized, and immersive digital humanities ecosystem through the following enhancements:", space_after=6)
    
    add_bullet(doc, "Integrating WebGL and Three.js rendering pipelines to allow users to rotate, examine, and inspect 3D photogrammetric scans of archaeological artifacts in 360 degrees.", bold_prefix="1. 3D Photogrammetry & Spatial Artifact Inspection: ")
    add_bullet(doc, "Deploying specialized, grounded LLM personas enabling students to hold historically accurate conversations with avatars of Emperor Ashoka, Aryabhata, or Akbar.", bold_prefix="2. Generative AI Historical Persona Chatbots: ")
    add_bullet(doc, "Leveraging WebXR on mobile devices to project life-sized architectural monuments (e.g. Sanchi Stupa, Konark Wheel) directly into physical classrooms.", bold_prefix="3. Augmented Reality (AR) Monument Projections: ")
    add_bullet(doc, "Synthesizing authentic voiceovers across Hindi, Tamil, Bengali, Marathi, Telugu, and Sanskrit using neural text-to-speech models.", bold_prefix="4. Multilingual Curatorial Audio Guides: ")
    add_bullet(doc, "Enabling verified historians and university departments to submit peer-reviewed curatorial annotations via cryptographic digital signatures.", bold_prefix="5. Decentralized Academic Peer Review: ")

def build_p66(doc):
    # Page 66: Chapter 15: Conclusion
    add_heading_1(doc, "15. CONCLUSION")
    add_p(doc, 
          "The 'History of India' project successfully addresses the critical limitations of contemporary digital heritage "
          "platforms: fragmented archives, non-interactive encyclopedic walls of text, broken media links, and the complete "
          "absence of acoustic atmosphere. By adopting a decoupled multi-tier architecture, the platform synergistically "
          "unifies modern software engineering with rigorous historical curation.", space_after=6)

    add_p(doc, 
          "Key engineering accomplishments of this project include:", space_after=4)
    add_bullet(doc, "Developing a high-performance Python 3.12 standard library server with zero third-party framework overhead, delivering sub-second response times.", bold_prefix="Lightweight Backend Rigor: ")
    add_bullet(doc, "Engineering a dual-mode chronological timeline that intuitively conveys civilizational continuity across five millennia.", bold_prefix="Chronological Fluency: ")
    add_bullet(doc, "Curating 22 high-resolution museum relics with verified provenance and full-screen inspection capabilities.", bold_prefix="Curatorial Excellence: ")
    add_bullet(doc, "Synthesizing an authentic Indian classical Tanpura drone procedurally via native browser Web Audio API oscillator nodes.", bold_prefix="Procedural Acoustic Innovation: ")
    add_bullet(doc, "Engaging users through gamified knowledge assessment quizzes and instant multi-index search filtering.", bold_prefix="Pedagogical Impact: ")

    add_p(doc, 
          "In conclusion, the platform establishes a reproducible, production-ready benchmark for next-generation digital "
          "humanities applications, demonstrating how computer science can honor, preserve, and illuminate India's "
          "extraordinary civilizational legacy for generations to come.", space_after=4)

def build_p67(doc):
    # Page 67: Chapter 16: Glossary
    add_heading_1(doc, "16. GLOSSARY")
    add_p(doc, "Alphabetical reference of academic, historical, and software engineering terminology utilized in this report:", space_after=4)

    tbl_glo = doc.add_table(rows=1, cols=2)
    tbl_glo.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl_glo.rows[0].cells
    hdr[0].width = Inches(2.0)
    hdr[1].width = Inches(4.6)
    for i, h in enumerate(["Term", "Definition & Context"]):
        set_cell_background(hdr[i], "1E293B")
        set_cell_margins(hdr[i], top=40, bottom=40, left=80, right=80)
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    glossary = [
        ("Digital Humanities", "An interdisciplinary field bridging computational methods and humanities scholarship."),
        ("Web Audio API", "A high-level W3C JavaScript API for processing and synthesizing audio in web applications."),
        ("OscillatorNode", "An audio processing module in Web Audio API generating periodic mathematical waveforms."),
        ("GainNode", "A Web Audio module controlling the amplitude/volume envelope of an audio signal."),
        ("Sony Darkroom", "A high-contrast dark aesthetic featuring slate and gold accents inspired by pro equipment."),
        ("Glassmorphism", "A UI design trend utilizing translucent frosted-glass backgrounds with backdrop-filter blur."),
        ("Dual-Mode Timeline", "A navigation interface allowing switching between horizontal carousel and vertical feed."),
        ("Indus Valley", "Bronze Age urban civilization (c. 2500 BCE) famous for planned brick cities and sanitation."),
        ("Maurya Empire", "Ancient pan-Indian empire (322–185 BCE) founded by Chandragupta Maurya and ruled by Ashoka."),
        ("Gupta Golden Age", "Classical Indian epoch (319–550 CE) marked by breakthroughs in math, science, and art."),
        ("MIME Type", "Standard identifying the format of files transmitted across the Internet (e.g. text/css)."),
        ("REST API", "Representational State Transfer architectural style for lightweight web services."),
        ("Fisher-Yates Shuffle", "An algorithm for generating a random permutation of a finite sequence in O(n) time.")
    ]
    for trm, dfn in glossary:
        row = tbl_glo.add_row().cells
        row[0].width = Inches(2.0)
        row[1].width = Inches(4.6)
        bg = "F8FAFC" if (len(tbl_glo.rows) % 2 == 0) else "FFFFFF"
        for i, val in enumerate([trm, dfn]):
            set_cell_background(row[i], bg)
            set_cell_margins(row[i], top=22, bottom=22, left=60, right=60)
            p = row[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)

def build_p68(doc):
    # Page 68: Chapter 17: References
    add_heading_1(doc, "17. REFERENCES")
    add_p(doc, "Scholarly monographs, W3C standards, historical catalogs, and software engineering treatises:", space_after=4)

    refs = [
        "1. Archaeological Survey of India (ASI). (2025). National Mission on Monuments and Antiquities. Retrieved from https://asi.nic.in",
        "2. National Museum, New Delhi. (2024). Digital Collection & Catalog of Antiquities. Retrieved from https://nationalmuseumindia.gov.in",
        "3. Thapar, Romila. (2002). The Penguin History of Early India: From the Origins to AD 1300. Penguin Books.",
        "4. Majumdar, R.C. (1977). Ancient India (8th ed.). Motilal Banarsidass Publishers.",
        "5. Eaton, Richard M. (2019). India in the Persianate Age: 1000–1765. Allen Lane.",
        "6. Chandra, Bipan. (1989). India's Struggle for Independence. Penguin Books India.",
        "7. Keay, John. (2000). India: A History. HarperCollins Publishers.",
        "8. Python Software Foundation. (2026). Python 3.12 Standard Library Reference: http.server and mimetypes. https://www.python.org/doc",
        "9. W3C Web Audio Working Group. (2024). Web Audio API W3C Recommendation. World Wide Web Consortium. https://www.w3.org/TR/webaudio",
        "10. Flanagan, David. (2020). JavaScript: The Definitive Guide (7th ed.). O'Reilly Media.",
        "11. Mozilla Developer Network (MDN). (2026). Web APIs: AudioContext, OscillatorNode, and SVG. https://developer.mozilla.org",
        "12. Fielding, Roy Thomas. (2000). Architectural Styles and the Design of Network-based Software Architectures (Doctoral dissertation).",
        "13. Gamma, Erich, Helm, Richard, Johnson, Ralph, & Vlissides, John. (1994). Design Patterns: Elements of Reusable Object-Oriented Software. Addison-Wesley.",
        "14. W3C SVG Working Group. (2018). Scalable Vector Graphics (SVG) 2 Specification. https://www.w3.org/TR/SVG2",
        "15. World Wide Web Consortium. (2018). Web Content Accessibility Guidelines (WCAG) 2.1. https://www.w3.org/TR/WCAG21",
        "16. British Museum. (2025). South Asia Antiquities Catalog and Provenance Records. https://www.britishmuseum.org/collection",
        "17. Victoria and Albert Museum. (2024). Arts of South Asia and the Mughal Court Collection. https://www.vam.ac.uk",
        "18. Childe, V. Gordon. (1950). The Urban Revolution. The Town Planning Review, 21(1), 3–17.",
        "19. Pressman, Roger S., & Maxim, Bruce R. (2020). Software Engineering: A Practitioner's Approach (9th ed.). McGraw-Hill Education.",
        "20. University of Mumbai. (2026). Guidelines for Final Year B.Sc Computer Science Project Work (Sem V & VI). University of Mumbai."
    ]

    for rf in refs:
        add_p(doc, rf, space_after=3)

print("pages_56_to_68 loaded.")
