import docx

doc = docx.Document(r'C:\Users\RAUNAK JAISWAL\OneDrive\Desktop\History_of_India_Project_Report_Raunak_Jaiswal.docx')
print('Number of paragraphs:', len(doc.paragraphs))
print('Number of tables:', len(doc.tables))

page_breaks = 0
for p in doc.paragraphs:
    for run in p.runs:
        if 'w:type="page"' in run._r.xml:
            page_breaks += 1

print('Explicit page breaks found:', page_breaks)
print('Embedded images (inline_shapes):', len(doc.inline_shapes))

sec = doc.sections[0]
xml_str = sec._sectPr.xml
print('Page borders in XML:', 'pgBorders' in xml_str)

ftr_xml = sec.footer.paragraphs[0]._p.xml
print('Footer has PAGE fldSimple:', 'PAGE' in ftr_xml)
print('Footer has of 68:', 'of 68' in ftr_xml)

# Check search for student name
full_text = " ".join([p.text for p in doc.paragraphs])
print('Contains Raunak Dilip Jaiswal:', 'Raunak Dilip Jaiswal' in full_text)
print('Contains V.K. Krishna Menon College:', 'V.K. Krishna Menon College' in full_text)
print('Contains MR. Pritam Karanjkar:', 'Pritam Karanjkar' in full_text)
print('Contains History of India:', 'History of India' in full_text)
print('Contains AI stock analysis (should be False):', 'AI stock analysis' in full_text or 'AI Stock Analysis' in full_text)
