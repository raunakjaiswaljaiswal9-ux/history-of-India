import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from report_lib import add_heading_1, add_heading_2, add_p, add_bullet, set_cell_background, set_cell_margins

def build_p1(doc):
    # Cover / Title Page
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(12)
    r = p_top.add_run("A\nPROJECT REPORT\nOn")
    r.font.name = 'Calibri'
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(24)
    r_t = p_title.add_run("HISTORY OF INDIA\nAn Interactive Digital Archive & Full-Stack Web Platform")
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(20)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(12)
    p_sub.paragraph_format.space_after = Pt(8)
    r_s = p_sub.add_run("SUBMITTED BY\n")
    r_s.font.name = 'Calibri'
    r_s.font.size = Pt(12)
    r_s.font.color.rgb = RGBColor(71, 85, 105)
    
    r_name = p_sub.add_run("Mr. Raunak Dilip Jaiswal")
    r_name.font.name = 'Calibri'
    r_name.font.size = Pt(16)
    r_name.bold = True
    r_name.font.color.rgb = RGBColor(15, 23, 42)

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(12)
    p_deg.paragraph_format.space_after = Pt(12)
    r_d = p_deg.add_run("In partial fulfilment for the award of the degree\nOf\nBACHELOR OF SCIENCE\nIn\nCOMPUTER SCIENCE")
    r_d.font.name = 'Calibri'
    r_d.font.size = Pt(13)
    r_d.bold = True
    r_d.font.color.rgb = RGBColor(30, 41, 59)

    p_gui = doc.add_paragraph()
    p_gui.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_gui.paragraph_format.space_before = Pt(10)
    p_gui.paragraph_format.space_after = Pt(6)
    r_g = p_gui.add_run("Under the guidance of\nMR. Pritam Karanjkar\nDepartment of Computer Science")
    r_g.font.name = 'Calibri'
    r_g.font.size = Pt(12)
    r_g.bold = True
    r_g.font.color.rgb = RGBColor(51, 65, 85)

    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col.paragraph_format.space_before = Pt(14)
    p_col.paragraph_format.space_after = Pt(6)
    r_c = p_col.add_run(
        "Bhandup Educational Society's\n"
        "V.K. KRISHNA MENON COLLEGE OF COMMERCE AND ECONOMICS\n"
        "AND SHARAD SHANKAR DIGHE COLLEGE OF SCIENCE AND P.A. MENON COLLEGE OF MANAGEMENT\n"
        "[AFFILIATED TO UNIVERSITY OF MUMBAI] BHANDUP EAST, MUMBAI-400042\n"
        "(SEM V)\n(2026-27)"
    )
    r_c.font.name = 'Calibri'
    r_c.font.size = Pt(11)
    r_c.bold = True
    r_c.font.color.rgb = RGBColor(15, 23, 42)

def build_p2(doc):
    # Certificate Page - Official Institutional Header
    p_soc = doc.add_paragraph()
    p_soc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_soc.paragraph_format.space_before = Pt(4)
    p_soc.paragraph_format.space_after = Pt(2)
    r_soc = p_soc.add_run("Bhandup Educational Society's")
    r_soc.font.name = 'Calibri'
    r_soc.font.size = Pt(14)
    r_soc.bold = True
    r_soc.font.color.rgb = RGBColor(15, 23, 42)

    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col.paragraph_format.space_before = Pt(2)
    p_col.paragraph_format.space_after = Pt(2)
    p_col.paragraph_format.line_spacing = 1.15
    r_col = p_col.add_run(
        "V.K. KRISHNA MENON COLLEGE OF COMMERCE & ECONOMICS AND\n"
        "SHARAD SHANKAR DIGHE COLLEGE OF SCIENCE AND P.A. MENON\n"
        "COLLEGE OF MANAGEMENT"
    )
    r_col.font.name = 'Calibri'
    r_col.font.size = Pt(12)
    r_col.bold = True
    r_col.font.color.rgb = RGBColor(15, 23, 42)

    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_before = Pt(2)
    p_aff.paragraph_format.space_after = Pt(2)
    r_aff = p_aff.add_run("[AFFILIATED TO UNIVERSITY OF MUMBAI]\nBHANDUP EAST, MUMBAI-400042")
    r_aff.font.name = 'Calibri'
    r_aff.font.size = Pt(11)
    r_aff.bold = True
    r_aff.font.color.rgb = RGBColor(30, 41, 59)

    p_dep = doc.add_paragraph()
    p_dep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dep.paragraph_format.space_before = Pt(4)
    p_dep.paragraph_format.space_after = Pt(6)
    r_dep = p_dep.add_run("DEPARTMENT OF COMPUTER SCIENCE")
    r_dep.font.name = 'Calibri'
    r_dep.font.size = Pt(13)
    r_dep.bold = True
    r_dep.font.color.rgb = RGBColor(15, 23, 42)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(12)
    r_line = p_line.add_run("____________________________________________________________")
    r_line.font.name = 'Calibri'
    r_line.font.size = Pt(10)
    r_line.bold = True
    r_line.font.color.rgb = RGBColor(148, 163, 184)

    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert.paragraph_format.space_before = Pt(4)
    p_cert.paragraph_format.space_after = Pt(14)
    r_ct = p_cert.add_run("CERTIFICATE")
    r_ct.font.name = 'Calibri'
    r_ct.font.size = Pt(18)
    r_ct.bold = True
    r_ct.font.color.rgb = RGBColor(15, 23, 42)

    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.space_before = Pt(8)
    p_body.paragraph_format.space_after = Pt(18)
    p_body.paragraph_format.line_spacing = 1.3
    r_b = p_body.add_run(
        "This is to certify that Mr. Raunak Dilip Jaiswal of T.Y. B.Sc (Sem V) class has "
        "satisfactorily completed the Project entitled \"History of India - An Interactive Digital "
        "Archive & Full-Stack Web Platform\" to be submitted in the partial fulfilment for the award "
        "of Bachelor of Science in Computer Science during the academic year 2026-2027."
    )
    r_b.font.name = 'Calibri'
    r_b.font.size = Pt(12)
    r_b.font.color.rgb = RGBColor(30, 41, 59)

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(10)
    p_date.paragraph_format.space_after = Pt(20)
    r_d = p_date.add_run("Date of Submission: ____________________")
    r_d.font.name = 'Calibri'
    r_d.font.size = Pt(11.5)
    r_d.bold = True

    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    c00 = tbl.cell(0, 0)
    c00.width = Inches(3.3)
    p00 = c00.paragraphs[0]
    p00.add_run("_________________________\nProject Guide\n(MR. Pritam Karanjkar)").font.name = 'Calibri'
    p00.runs[0].bold = True
    p00.runs[0].font.size = Pt(11)

    c01 = tbl.cell(0, 1)
    c01.width = Inches(3.3)
    p01 = c01.paragraphs[0]
    p01.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p01.add_run("_________________________\nHead / Incharge\nDepartment of Computer Science").font.name = 'Calibri'
    p01.runs[0].bold = True
    p01.runs[0].font.size = Pt(11)

    c10 = tbl.cell(1, 0)
    c10.width = Inches(3.3)
    p10 = c10.paragraphs[0]
    p10.paragraph_format.space_before = Pt(24)
    p10.add_run("\nCollege Seal").font.name = 'Calibri'
    p10.runs[0].bold = True
    p10.runs[0].font.size = Pt(11)

    c11 = tbl.cell(1, 1)
    c11.width = Inches(3.3)
    p11 = c11.paragraphs[0]
    p11.paragraph_format.space_before = Pt(24)
    p11.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p11.add_run("\n_________________________\nSignature of Examiner").font.name = 'Calibri'
    p11.runs[0].bold = True
    p11.runs[0].font.size = Pt(11)

def build_p3(doc):
    # Declaration Page
    p_dec = doc.add_paragraph()
    p_dec.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dec.paragraph_format.space_before = Pt(24)
    p_dec.paragraph_format.space_after = Pt(24)
    r_dc = p_dec.add_run("DECLARATION")
    r_dc.font.name = 'Calibri'
    r_dc.font.size = Pt(18)
    r_dc.bold = True
    r_dc.font.color.rgb = RGBColor(15, 23, 42)

    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.space_before = Pt(12)
    p_body.paragraph_format.space_after = Pt(20)
    p_body.paragraph_format.line_spacing = 1.35
    r_b = p_body.add_run(
        "I, Mr. Raunak Dilip Jaiswal hereby declare that the project entitled "
        "\"History of India - An Interactive Digital Archive & Full-Stack Web Platform\" submitted in "
        "the partial fulfilment for the award of Bachelor of Science in Computer Science during the "
        "academic year 2026-2027 is my original work and the project has not formed the basis for the "
        "award of any degree, associateship, fellowship or any other similar titles."
    )
    r_b.font.name = 'Calibri'
    r_b.font.size = Pt(12)
    r_b.font.color.rgb = RGBColor(30, 41, 59)

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(50)
    p_sig.paragraph_format.space_after = Pt(12)
    r_s = p_sig.add_run("Signature of the Student: _________________________\nMr. Raunak Dilip Jaiswal")
    r_s.font.name = 'Calibri'
    r_s.font.size = Pt(12)
    r_s.bold = True

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(24)
    p_meta.paragraph_format.space_after = Pt(8)
    r_m = p_meta.add_run("Place: Mumbai\n\nDate:  _________________________")
    r_m.font.name = 'Calibri'
    r_m.font.size = Pt(12)
    r_m.bold = True

def build_p4(doc):
    # Acknowledgement Page
    p_ack = doc.add_paragraph()
    p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack.paragraph_format.space_before = Pt(18)
    p_ack.paragraph_format.space_after = Pt(18)
    r_ak = p_ack.add_run("ACKNOWLEDGEMENT")
    r_ak.font.name = 'Calibri'
    r_ak.font.size = Pt(18)
    r_ak.bold = True
    r_ak.font.color.rgb = RGBColor(15, 23, 42)

    add_p(doc, 
          "I would like to extend my deepest gratitude to Mr. Prashant Rajan, our respected Principal, "
          "and all staff of V.K. Krishna Menon College of Commerce & Economics and Sharad Shankar Dighe College "
          "of Science for providing us moral support, an enriching learning environment, and the much-needed "
          "inspiration to complete this project on time.", space_after=10)

    add_p(doc, 
          "I wish to extend my sincerest gratitude and heartfelt thanks to my esteemed project guide "
          "MR. Pritam Karanjkar for giving his generous support, profound academic insights, technical "
          "direction, and patient mentorship throughout the entire research and development lifecycle of this project.", space_after=10)

    add_p(doc, 
          "I would also like to convey my special thanks to the Management and Department of Computer Science "
          "for providing the requisite laboratories, digital infrastructure, and computing facilities that enabled "
          "the seamless execution and curatorial refinement of this web platform.", space_after=10)

    add_p(doc, 
          "I am extremely grateful to the University of Mumbai for prescribing this comprehensive project work "
          "as an essential component of the final year Bachelor of Science in Computer Science curriculum, allowing "
          "me to apply theoretical software engineering principles to real-world cultural heritage informatics.", space_after=10)

    add_p(doc, 
          "Finally, I express my sincere appreciation to my parents, family members, and fellow classmates who "
          "have directly or indirectly supported, encouraged, and assisted me in successfully bringing this "
          "ambitious project to fruition.", space_after=10)

def build_p5(doc):
    # Table of Contents Page
    p_toc = doc.add_paragraph()
    p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc.paragraph_format.space_before = Pt(8)
    p_toc.paragraph_format.space_after = Pt(8)
    r_tc = p_toc.add_run("TABLE OF CONTENT")
    r_tc.font.name = 'Calibri'
    r_tc.font.size = Pt(16)
    r_tc.bold = True
    r_tc.font.color.rgb = RGBColor(15, 23, 42)

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    hdr[0].width = Inches(0.8)
    hdr[1].width = Inches(4.7)
    hdr[2].width = Inches(1.1)
    
    headers = ["Sr.No", "Contents", "Pg. No."]
    for i, h in enumerate(headers):
        set_cell_background(hdr[i], "1E293B")
        set_cell_margins(hdr[i], top=60, bottom=60, left=80, right=80)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(10.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    toc_items = [
        ("1", "Project Proposal", "6-8"),
        ("2", "Plagiarism Report", "9-10"),
        ("3", "Introduction", "11"),
        ("4", "Objective", "12-13"),
        ("5", "Scope", "14-16"),
        ("6", "Methodology", "17-19"),
        ("7", "Requirement Specification", "20-23"),
        ("8", "Feasibility Study", "24-27"),
        ("9", "Gantt Chart", "28"),
        ("10", "System Design", "29-35"),
        ("10.1", "ER Diagram", "36-37"),
        ("10.2", "Class Diagram", "38"),
        ("10.3", "Activity Diagram", "39"),
        ("10.4", "Use Case Diagram", "40"),
        ("11", "Code Implementation", "41-55"),
        ("12", "Database", "56-57"),
        ("13", "Screenshots", "58-64"),
        ("14", "Future Scope", "65"),
        ("15", "Conclusion", "66"),
        ("16", "Glossary", "67"),
        ("17", "References", "68")
    ]

    for sr, cnt, pg in toc_items:
        row = tbl.add_row().cells
        row[0].width = Inches(0.8)
        row[1].width = Inches(4.7)
        row[2].width = Inches(1.1)
        
        bg_col = "F8FAFC" if (len(tbl.rows) % 2 == 0) else "FFFFFF"
        for i, text in enumerate([sr, cnt, pg]):
            set_cell_background(row[i], bg_col)
            set_cell_margins(row[i], top=35, bottom=35, left=70, right=70)
            p = row[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(10)
            if sr in ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13", "14", "15", "16", "17"]:
                r.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)

print("pages_01_to_05 loaded.")
