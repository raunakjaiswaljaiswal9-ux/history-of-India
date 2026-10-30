#!/usr/bin/env python3
"""
Master Builder for 68-Page University of Mumbai Project Report
Candidate: Mr. Raunak Dilip Jaiswal
Project: History of India - An Interactive Digital Archive & Full-Stack Web Platform
College: V.K. Krishna Menon College of Commerce & Economics, Mumbai-400042
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from pages_01_to_05 import build_p1, build_p2, build_p3, build_p4, build_p5
from pages_06_to_10 import build_p6, build_p7, build_p8, build_p9, build_p10
from pages_11_to_20 import (
    build_p11, build_p12, build_p13, build_p14, build_p15,
    build_p16, build_p17, build_p18, build_p19, build_p20
)
from pages_21_to_30 import (
    build_p21, build_p22, build_p23, build_p24, build_p25,
    build_p26, build_p27, build_p28, build_p29, build_p30
)
from pages_31_to_40 import (
    build_p31, build_p32, build_p33, build_p34, build_p35,
    build_p36, build_p37, build_p38, build_p39, build_p40
)
from pages_41_to_55 import (
    build_p41, build_p42, build_p43, build_p44, build_p45,
    build_p46, build_p47, build_p48, build_p49, build_p50,
    build_p51, build_p52, build_p53, build_p54, build_p55
)
from pages_56_to_68 import (
    build_p56, build_p57, build_p58, build_p59, build_p60,
    build_p61, build_p62, build_p63, build_p64, build_p65,
    build_p66, build_p67, build_p68
)

PAGE_BUILDERS = [
    build_p1, build_p2, build_p3, build_p4, build_p5,
    build_p6, build_p7, build_p8, build_p9, build_p10,
    build_p11, build_p12, build_p13, build_p14, build_p15,
    build_p16, build_p17, build_p18, build_p19, build_p20,
    build_p21, build_p22, build_p23, build_p24, build_p25,
    build_p26, build_p27, build_p28, build_p29, build_p30,
    build_p31, build_p32, build_p33, build_p34, build_p35,
    build_p36, build_p37, build_p38, build_p39, build_p40,
    build_p41, build_p42, build_p43, build_p44, build_p45,
    build_p46, build_p47, build_p48, build_p49, build_p50,
    build_p51, build_p52, build_p53, build_p54, build_p55,
    build_p56, build_p57, build_p58, build_p59, build_p60,
    build_p61, build_p62, build_p63, build_p64, build_p65,
    build_p66, build_p67, build_p68
]

def generate_report():
    print(f"[*] Starting 68-page report generation. Total page builders: {len(PAGE_BUILDERS)}...")
    doc = Document()

    # Section & Page Setup: A4 format with 0.75 in margins
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    # 1. Add Page Border (single black line with 24pt offset, identical to reference report)
    borders = parse_xml(f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
                        f'<w:top w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
                        f'<w:left w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
                        f'<w:bottom w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
                        f'<w:right w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
                        f'</w:pgBorders>')
    sec._sectPr.append(borders)

    # 2. Header and Footer Configuration
    sec.different_first_page_header_footer = True

    # Word Header on Page 2 (Certificate) and subsequent pages
    header = sec.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.paragraph_format.space_after = Pt(0)
    r_hdr = p_hdr.add_run("V.K. KRISHNA MENON COLLEGE | DEPARTMENT OF COMPUTER SCIENCE")
    r_hdr.font.name = 'Calibri'
    r_hdr.font.size = Pt(9)
    r_hdr.font.color.rgb = RGBColor(100, 116, 139)

    # Right-Aligned Footer: 'Page X of 68' in Calibri 11pt on all pages
    for ftr in [sec.footer, sec.first_page_footer]:
        p_ftr = ftr.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ftr.paragraph_format.space_after = Pt(0)
        
        r1 = p_ftr.add_run("Page ")
        r1.font.name = 'Calibri'
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(71, 85, 105)

        fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        p_ftr._p.append(fld)

        r2 = p_ftr.add_run(" of 68")
        r2.font.name = 'Calibri'
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # 3. Build each page sequentially with exact page breaks
    for idx, builder in enumerate(PAGE_BUILDERS):
        page_num = idx + 1
        print(f"  [+] Building Page {page_num} of 68: {builder.__name__}...")
        builder(doc)
        if page_num < 68:
            doc.add_page_break()

    # 4. Save output documents
    workspace_out = r"C:\Users\RAUNAK JAISWAL\OneDrive\Desktop\History of India\History_of_India_Project_Report.docx"
    desktop_out = r"C:\Users\RAUNAK JAISWAL\OneDrive\Desktop\History_of_India_Project_Report_Raunak_Jaiswal.docx"

    print(f"[*] Saving master document to {workspace_out}...")
    doc.save(workspace_out)
    print(f"[+] Workspace copy saved successfully ({os.path.getsize(workspace_out)} bytes).")

    print(f"[*] Copying final report to Desktop at {desktop_out}...")
    shutil.copy2(workspace_out, desktop_out)
    print(f"[+] Desktop copy saved successfully ({os.path.getsize(desktop_out)} bytes).")
    print("[SUCCESS] All 68 pages compiled and saved flawlessly!")

if __name__ == "__main__":
    generate_report()
